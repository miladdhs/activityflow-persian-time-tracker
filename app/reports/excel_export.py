"""Excel export with right-to-left sheets."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.core.errors import StorageError
from app.reports.tables import RECORD_HEADERS, group_headers, group_row, record_row, summary_rows
from app.services.activity_service import ActivityRow
from app.services.report_service import ReportResult

FONT = "Vazirmatn"
HEADER_FILL = PatternFill("solid", fgColor="2A1F45")
HEADER_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(name=FONT, size=11)
TITLE_FONT = Font(name=FONT, bold=True, size=16, color="5B3CC4")
LABEL_FONT = Font(name=FONT, bold=True, size=11, color="4A4560")
THIN = Side(style="thin", color="DDD8EA")
RIGHT = Alignment(horizontal="right", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center")


def _table(ws, headers: list[str], rows: list[list], widths: list[int], start_row: int = 1) -> None:
    ws.sheet_view.rightToLeft = True
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=start_row, column=col, value=header)
        cell.font, cell.fill, cell.alignment = HEADER_FONT, HEADER_FILL, CENTER
    for r, values in enumerate(rows, start_row + 1):
        for c, value in enumerate(values, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.font = BODY_FONT
            cell.border = Border(bottom=THIN)
            cell.alignment = CENTER if isinstance(value, (int, float)) else RIGHT
    for col, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[start_row].height = 26
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(headers))}{start_row + len(rows)}"


def export_report_xlsx(
    result: ReportResult, records: list[ActivityRow], meta: list[tuple[str, str]], path: Path
) -> Path:
    wb = Workbook()

    ws = wb.active
    ws.title = "خلاصه"
    ws.sheet_view.rightToLeft = True
    ws["A1"] = "گزارش ActivityFlow"
    ws["A1"].font = TITLE_FONT
    row = 3
    for label, value in [*meta, ("", ""), *summary_rows(result)]:
        if label:
            ws.cell(row=row, column=1, value=label).font = LABEL_FONT
            cell = ws.cell(row=row, column=2, value=value)
            cell.font = BODY_FONT
            cell.alignment = RIGHT
        row += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 46

    _table(
        wb.create_sheet("گروه‌ها"),
        group_headers(result.group_by),
        [group_row(g) for g in result.groups],
        [34, 12, 22, 18, 24, 16, 34, 16],
    )
    _table(
        wb.create_sheet("رکوردها"),
        RECORD_HEADERS,
        [record_row(r) for r in records],
        [14, 12, 32, 24, 20, 12, 12, 22, 16, 12, 12, 40],
    )

    try:
        wb.save(path)
    except OSError as exc:
        raise StorageError("ذخیره فایل Excel با خطا مواجه شد. ممکن است فایل در برنامه دیگری باز باشد.") from exc
    return Path(path)
