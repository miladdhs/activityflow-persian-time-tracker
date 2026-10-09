from datetime import date, time, timedelta
from decimal import Decimal

import pytest

from app.core.errors import ConflictError, ValidationError


def _units():
    from app.services import lookup_service

    return {u.name: u.id for u in lookup_service.list_units()}


def _add(title, day=None, minutes=60, project=None, quantity=None, unit=None, category=None):
    from app.services.activity_service import ActivityInput, create_activity

    return create_activity(
        ActivityInput(
            title=title,
            activity_date=day or date.today(),
            project_id=project,
            category_id=category,
            duration_minutes=minutes,
            quantity=quantity,
            unit_id=unit,
        )
    )


def test_seed_data(db):
    from app.services import lookup_service

    assert len(lookup_service.list_categories()) == 8
    assert "عدد" in _units()


def test_same_name_activities_share_one_title(db):
    from app.services import lookup_service

    _add("کلیپ کاغذ")
    _add("كليپ  كاغذ")
    _add("کلیپ کاغذ تبلیغاتی")
    titles = {t.name: t.usage for t in lookup_service.list_titles()}
    assert titles == {"کلیپ کاغذ": 2, "کلیپ کاغذ تبلیغاتی": 1}
    assert lookup_service.similar_titles("کلیپ کاغذ")


def test_validation_rules(db):
    from app.services.activity_service import ActivityInput, create_activity

    today = date.today()
    cases = [
        (ActivityInput(title="", activity_date=today, duration_minutes=10), "title"),
        (ActivityInput(title="x", activity_date=today + timedelta(days=1), duration_minutes=10), "date"),
        (ActivityInput(title="x", activity_date=today, duration_minutes=0), "duration"),
        (ActivityInput(title="x", activity_date=today, duration_minutes=2000), "duration"),
        (ActivityInput(title="x", activity_date=today, time_mode="range", start_time=time(10), end_time=time(9)), "end"),
        (ActivityInput(title="x", activity_date=today, duration_minutes=10, quantity="-2", unit_id=1), "quantity"),
        (ActivityInput(title="x", activity_date=today, duration_minutes=10, quantity="1.2345", unit_id=1), "quantity"),
        (ActivityInput(title="x", activity_date=today, duration_minutes=10, quantity="5"), "unit"),
    ]
    for inp, field in cases:
        with pytest.raises(ValidationError) as err:
            create_activity(inp)
        assert err.value.field == field


def test_time_range_computes_duration(db):
    from app.services.activity_service import ActivityInput, create_activity, get_activity

    aid = create_activity(
        ActivityInput(title="جلسه تیم", activity_date=date.today(), time_mode="range",
                      start_time=time(9, 15), end_time=time(11, 0))
    )
    row = get_activity(aid)
    assert row.duration_minutes == 105
    assert row.time_mode == "range"


def test_quantity_only_activity(db):
    from app.services.activity_service import get_activity

    aid = _add("عکاسی محصول", minutes=0, quantity="۱۲٫۵", unit=_units()["عکس"])
    row = get_activity(aid)
    assert row.quantity == Decimal("12.5") and row.unit_name == "عکس"


def test_report_groups_across_projects_and_keeps_units_apart(db):
    from app.services import project_service, report_service
    from app.services.activity_service import ActivityFilter
    from app.services.project_service import ProjectInput

    a = project_service.create_project(ProjectInput(name="پروژه A"))
    b = project_service.create_project(ProjectInput(name="پروژه B"))
    units = _units()
    _add("کلیپ کاغذ", minutes=30, project=a, quantity="40", unit=units["کلیپ"])
    _add("کلیپ کاغذ", minutes=45, project=b, quantity="60", unit=units["کلیپ"])
    _add("کلیپ کاغذ", minutes=15, project=b, quantity="2", unit=units["کیلوگرم"])
    _add("جلسه", minutes=60, project=a)

    today = date.today()
    f = ActivityFilter(date_from=today - timedelta(days=3), date_to=today)
    result = report_service.build_report(f, "title")
    assert result.count == 4 and result.total_minutes == 150
    clip = next(g for g in result.groups if g.name == "کلیپ کاغذ")
    assert clip.count == 3 and clip.minutes == 90
    assert dict(clip.quantities) == {"کلیپ": Decimal(100), "کیلوگرم": Decimal(2)}
    records = report_service.group_records(f, "title", clip.key)
    assert len(records) == 3
    assert sum(r.duration_minutes for r in records) == clip.minutes

    only_a = report_service.build_report(ActivityFilter(date_from=f.date_from, date_to=today, project_id=a), "title")
    assert only_a.total_minutes == 90

    by_project = report_service.build_report(f, "project")
    assert {g.name: g.minutes for g in by_project.groups} == {"پروژه A": 90, "پروژه B": 60}


def test_report_updates_after_edit_and_delete(db):
    from app.services import activity_service, report_service
    from app.services.activity_service import ActivityFilter, ActivityInput

    aid = _add("طراحی لوگو", minutes=50)
    _add("طراحی لوگو", minutes=10)
    f = ActivityFilter(date_from=date.today(), date_to=date.today())
    assert report_service.build_report(f).total_minutes == 60

    activity_service.update_activity(aid, ActivityInput(title="طراحی بنر", activity_date=date.today(), duration_minutes=20))
    groups = {g.name: g.minutes for g in report_service.build_report(f).groups}
    assert groups == {"طراحی لوگو": 10, "طراحی بنر": 20}

    activity_service.delete_activity(aid)
    groups = {g.name: g.minutes for g in report_service.build_report(f).groups}
    assert groups == {"طراحی لوگو": 10}


def test_deleting_project_keeps_activities(db):
    from app.services import activity_service, project_service
    from app.services.project_service import ProjectInput

    pid = project_service.create_project(ProjectInput(name="موقت"))
    aid = _add("کار", project=pid)
    assert project_service.delete_project(pid) == 1
    assert activity_service.get_activity(aid).project_id is None


def test_archived_project_rejects_new_activities(db):
    from app.services import project_service
    from app.services.project_service import ProjectInput

    pid = project_service.create_project(ProjectInput(name="قدیمی"))
    project_service.set_status(pid, "archived")
    with pytest.raises(ValidationError):
        _add("کار", project=pid)


def test_duplicate_project_name(db):
    from app.services import project_service
    from app.services.project_service import ProjectInput

    project_service.create_project(ProjectInput(name="سایت فروشگاه"))
    with pytest.raises(ConflictError):
        project_service.create_project(ProjectInput(name="سايت  فروشگاه"))


def test_unit_in_use_cannot_be_deleted(db):
    from app.services import lookup_service

    unit = _units()["بسته"]
    _add("بسته‌بندی", quantity="3", unit=unit)
    with pytest.raises(ConflictError):
        lookup_service.delete_unit(unit)


def test_merge_titles(db):
    from app.services import lookup_service

    _add("کلیپ کاغذ")
    _add("کلیپ کاغذی")
    titles = {t.name: t.id for t in lookup_service.list_titles()}
    moved = lookup_service.merge_titles(titles["کلیپ کاغذی"], titles["کلیپ کاغذ"])
    assert moved == 1
    assert [(t.name, t.usage) for t in lookup_service.list_titles()] == [("کلیپ کاغذ", 2)]


def test_pagination_and_search(db):
    from app.services.activity_service import ActivityFilter, query_activities

    for i in range(30):
        _add(f"کار شماره {i}", day=date.today() - timedelta(days=i % 5), minutes=10 + i)
    page = query_activities(ActivityFilter(), page=2, page_size=25)
    assert page.total == 30 and page.page_count == 2 and len(page.rows) == 5
    found = query_activities(ActivityFilter(search="شماره ۱۲"))
    assert found.total == 1 and found.rows[0].title == "کار شماره 12"
    longest = query_activities(ActivityFilter(), sort="duration_desc", page_size=1)
    assert longest.rows[0].duration_minutes == 39


def test_backup_and_restore(db):
    from app.services import backup_service, lookup_service

    _add("قبل از پشتیبان")
    backup = backup_service.create_backup()
    _add("بعد از پشتیبان")
    assert len(lookup_service.list_titles()) == 2
    safety = backup_service.restore_backup(backup)
    assert safety.exists()
    assert [t.name for t in lookup_service.list_titles()] == ["قبل از پشتیبان"]


def test_restore_rejects_invalid_file(db):
    from app.services import backup_service

    bad = db / "bad.sqlite3"
    bad.write_text("not a database", encoding="utf-8")
    with pytest.raises(ValidationError):
        backup_service.restore_backup(bad)


def test_exports(db, tmp_path):
    import csv

    from openpyxl import load_workbook

    from app.reports.csv_export import export_records_csv
    from app.reports.excel_export import export_report_xlsx
    from app.services import report_service
    from app.services.activity_service import ActivityFilter, list_activities

    _add("گزارش نویسی", minutes=90)
    f = ActivityFilter(date_from=date.today(), date_to=date.today())
    rows = list_activities(f)
    path = export_records_csv(rows, tmp_path / "a.csv")
    with open(path, encoding="utf-8-sig") as fh:
        data = list(csv.reader(fh))
    assert data[0][2] == "عنوان فعالیت" and data[1][2] == "گزارش نویسی" and data[1][8] == "90"

    result = report_service.build_report(f)
    xlsx = export_report_xlsx(result, rows, [("بازه", "امروز")], tmp_path / "r.xlsx")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["خلاصه", "گروه‌ها", "رکوردها"]
    assert wb["گروه‌ها"].sheet_view.rightToLeft


def test_dashboard(db):
    from app.services import stats_service

    _add("الف", minutes=30)
    _add("الف", day=date.today() - timedelta(days=1), minutes=20)
    data = stats_service.dashboard()
    assert data.today_minutes == 30 and data.yesterday_minutes == 20
    assert len(data.week) == 7 and data.week[-1][1] == 30
    assert data.frequent[0].title == "الف" and data.frequent[0].count == 2
